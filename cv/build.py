from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_COLOR_INDEX, WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import re

d=Document()
s=d.sections[0]
s.page_width=Cm(21); s.page_height=Cm(29.7)
s.left_margin=s.right_margin=Cm(1.7); s.top_margin=Cm(1.2); s.bottom_margin=Cm(1.1)
st=d.styles['Normal']; st.font.name='Calibri'; st.font.size=Pt(10)
st.element.rPr.rFonts.set(qn('w:eastAsia'),'Calibri')
st.paragraph_format.space_after=Pt(0); st.paragraph_format.space_before=Pt(0)

def runs(p,text,size=None,bold=False,color=None):
    # **gras** et [à compléter] surligné
    for part in re.split(r'(\*\*.+?\*\*|\[[^\]]+\])',text):
        if not part: continue
        if part.startswith('**'):
            r=p.add_run(part[2:-2]); r.bold=True
        elif part.startswith('['):
            r=p.add_run(part); r.font.highlight_color=WD_COLOR_INDEX.YELLOW; r.bold=bold
        else:
            r=p.add_run(part); r.bold=bold
        if size: r.font.size=Pt(size)
        if color: r.font.color.rgb=RGBColor(*color)
    return p

def line(text,size=10,bold=False,align=None,after=0,color=None):
    p=d.add_paragraph(); runs(p,text,size,bold,color)
    if align: p.alignment=align
    p.paragraph_format.space_after=Pt(after); return p

def heading(t):
    p=d.add_paragraph(); r=p.add_run(t.upper()); r.bold=True; r.font.size=Pt(10.5); r.font.color.rgb=RGBColor(0x1F,0x2A,0x44)
    p.paragraph_format.space_before=Pt(7); p.paragraph_format.space_after=Pt(2)
    pPr=p._p.get_or_add_pPr(); b=OxmlElement('w:pBdr'); bt=OxmlElement('w:bottom')
    for k,v in (('val','single'),('sz','6'),('space','1'),('color','1F2A44')): bt.set(qn('w:'+k),v)
    b.append(bt); pPr.append(b)

def job(title,org,date):
    p=d.add_paragraph(); p.paragraph_format.space_before=Pt(4)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(17.6),alignment=2)
    r=p.add_run(title); r.bold=True
    p.add_run(' | '+org)
    r=p.add_run('\t'+date); r.bold=True

def bullet(t):
    p=d.add_paragraph(style='List Bullet'); runs(p,t)
    p.paragraph_format.left_indent=Cm(0.55); p.paragraph_format.first_line_indent=Cm(-0.4)

line('Mehdi TALIL',20,True,after=0,color=(0x1F,0x2A,0x44))
line('[TITRE EXACT DE L\'OFFRE — ex. Stage Consumer Insights / Stage E-Merchandising / Stage Marketing Digital] — 6 mois, dès janvier 2027',11,True,after=1)
line('Paris · Mobile en France  |  +33 7 76 48 75 98  |  mehdi.talil26@gmail.com  |  Anglais courant',9.5,after=2)

heading('Profil')
line('Étudiant en **M2 Marketing Digital** (double diplomation **ENCG Settat / Brest Business School**, Grande École). **Études de marché terrain** (Sunergia) et **catalogue e-commerce** de produits de grande consommation (Bladna) : je transforme des données consommateurs en actions business concrètes.')

heading('Compétences clés')
line('**Études & Insights :** études quantitatives et qualitatives, enquêtes terrain, benchmark concurrentiel, tendances de consommation, NielsenIQ, SPSS, Power BI, Excel (tableaux croisés, Power Query)')
line('**E-commerce & Digital :** gestion de catalogue produit, fiches produits, WordPress, Shopify, SEO / GEO, Meta Ads, CRM & fidélisation, KPIs, automation (n8n)')
line('**Outils :** Pack Office 365 (Excel, PowerPoint, Word), Photoshop, Illustrator, Canva, CapCut')

heading('Expériences professionnelles')
job('Enquêteur terrain','Sunergia','Juin – Août 2026')
bullet('Garanti la fiabilité des **données consommateurs** collectées, mesuré par le respect des objectifs quantitatifs et des délais de collecte ([N] questionnaires, [N] % de l\'objectif), en administrant questionnaires et guides d\'entretien en face-à-face.')
bullet('Maintenu un haut niveau de qualité de donnée en autonomie, mesuré par la conduite d\'enquêtes sur [N] zones variées, en organisant seul mes déplacements et ma collecte.')
job('Assistant Marketing & Développement','Bladna (Eurobrands), Paris','Janv. – Avr. 2026')
bullet('Amélioré le **référencement SEO / GEO** et les KPIs du site e-commerce Bladna.fr, mesuré par [+N % de trafic / N pages indexées], en refondant l\'architecture **WordPress** et en rédigeant [N] fiches produits et articles optimisés.')
bullet('Visé **+25 % de rétention client**, mesuré par le taux de réachat et les avis collectés, en concevant une stratégie de fidélisation (cartes de remerciement, incentives avis).')
bullet('Identifié les attentes de 2 segments (diaspora marocaine, consommateurs européens), mesuré par les insights tirés des retours clients, en analysant les tendances d\'un marché de grande consommation (FMCG).')
bullet('Sourcé et négocié avec **10+ fournisseurs européens** de cosmétiques, mesuré par [N fournisseurs retenus], en menant benchmark et négociation.')
job('Chef de cellule Communication & infographie','Club 6Days, ENCG Settat','2023 – 2024')
bullet('Renforcé la visibilité des activités associatives, mesuré par [N abonnés / N vues], en pilotant la communication interne et digitale (publications, reels, stories, affiches sur Instagram et TikTok).')
job('Assistant Chargé de Clientèle','CIH Bank, Casablanca','Juil. 2024')
bullet('Fiabilisé les dossiers clients, mesuré par la qualité des données saisies et vérifiées, en participant à l\'instruction de dossiers de crédit (~30 clients/jour).')

heading('Formation')
for t,o,dt,sub in (
 ('Master Marketing Digital – double diplomation (Grande École)','Brest Business School','2026 – 2027','Distribution omnicanale · Relation client · Stratégie de communication'),
 ('Master 1 Marketing & Actions Commerciales','ENCG Settat','2022 – 2027','Études de marché · Analyse de données · Marketing stratégique'),
 ('Neuroscience du consommateur & Neuromarketing','Copenhagen Business School (en ligne)','Avr. 2025','')):
    job(t,o,dt)
    if sub: line(sub,9.5)

heading('Engagements & langues')
bullet('**Bureau exécutif, Club 6Days (2024 – 2025)** : coordonné **22+ actions** événementielles et humanitaires, mesuré par la mobilisation de **200+ membres** ; **Prix du Meilleur Trésorier 2025** pour une gestion budgétaire rigoureuse.')
bullet('**Langues :** Français (bilingue) · Anglais (courant) · Arabe (langue maternelle) · Espagnol (A2)')
d.save('CV_Mehdi_TALIL_Master_XYZ.docx')
